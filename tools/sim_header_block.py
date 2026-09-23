"""Walks the template header block exactly as LCL_EXCEL=>READ now does.

Reads the customer's own workbook, finds the heading line the way the reader
does, then applies LCL_UTIL=>HEADER_MATTER to the unbroken run of lines under
it and reports which lines would be passed over and which first reaches a
handler.
"""
import re, sys, collections
import openpyxl

# X is not a type letter - see the comment in LCL_UTIL=>HEADER_MATTER
TYPES = {'C','N','D','T','P','I','F','CHAR','NUMC','DATS','TIMS','CURR',
         'DEC','QUAN','UNIT','CUKY','LANG','RAW'}


def squash(x):
    return re.sub(r'[^A-Z0-9]', '', (x or '').upper())


def header_matter(cells, keycol):
    vals = [(i + 1, str(c).strip().upper()) for i, c in enumerate(cells)
            if c is not None and str(c).strip()]
    if len(vals) < 2:
        return False
    n = len(vals)
    ty = sum(1 for _, v in vals if v in TYPES)
    dg = sum(1 for _, v in vals if v.isdigit())
    mo = sum(1 for _, v in vals if v in ('M', 'O', 'M/O'))
    long_dig = any(v.isdigit() and len(v) > 3 for _, v in vals)
    c1 = dict(vals).get(1)
    c1t, c1d, c1m = (c1 in TYPES), bool(c1 and c1.isdigit()), c1 in ('M', 'O', 'M/O')
    slack = 2 if n >= 4 else 1
    if ty > 0 and ty >= n - slack and (c1t or ty < n):
        return True
    if not long_dig and dg > 0 and dg >= n - slack:
        return True
    if mo > 0 and mo >= n - slack:
        return True
    key = dict(vals).get(keycol)
    return bool(key and ' ' in key)


def hdr_map(src):
    m = collections.defaultdict(list)
    for x in re.finditer(r"\(\s*scen = '(R\d)'\s+col = (\d+)\s+hdr = '([^']*)'\s*\)", src):
        m[x.group(1)].append(x.group(3))
    return m


def key_cols(src):
    """Each handler's key column, read per class block so a class without a
    redefinition does not borrow the next one's number."""
    out = {}
    for m in re.finditer(r'CLASS (lcl_h_\w+) IMPLEMENTATION\.(.*?)\nENDCLASS\.',
                         src, re.S):
        k = re.search(r'METHOD lif_h~key_col\.\s+rv = (\d+)', m.group(2))
        if k:
            out[m.group(1)] = int(k.group(1))
    return out


VEND = open('src/zmms_bp_mass_upload.prog.abap', encoding='utf8').read()
HDR = hdr_map(VEND)
KEYS = key_cols(VEND)
DEFAULT_KEY = 2
SCEN = [('R1', 'Vendor creation for All CC', 'lcl_h_create'),
        ('R2', 'TDS upload', 'lcl_h_tds'), ('R3', 'TAN details', 'lcl_h_tan'),
        ('R4', 'BANK Key creation', 'lcl_h_bkey'),
        ('R5', 'Bank details update', 'lcl_h_bank'),
        ('R6', 'Vendor extension', 'lcl_h_ext'),
        ('R7', 'CIN details', 'lcl_h_cin'),
        ('R8', 'Patner function', 'lcl_h_pfn'),
        ('R9', 'Block_Unblocked', 'lcl_h_blk')]

book = sys.argv[1] if len(sys.argv) > 1 else 'Vendor LSMW with Template.xlsx'
wb = openpyxl.load_workbook(book, data_only=True)
bad = []
for scen, tab, cls in SCEN:
    ws = wb[tab]
    rows = {r[0].row: [c.value for c in r] for r in ws.iter_rows(max_row=min(ws.max_row, 30))}
    want = {squash(h) for h in HDR[scen]}
    best, hrow = 0, 0
    for r in sorted(rows)[:10]:
        sc = len(want & {squash(str(v)) for v in rows[r] if v is not None})
        if sc > best:
            best, hrow = sc, r
    keycol = KEYS.get(cls, DEFAULT_KEY)
    skipped, data = [], hrow
    while data + 1 in rows:
        cells = rows[data + 1]
        if not any(c is not None and str(c).strip() for c in cells):
            break
        if not header_matter(cells, keycol):
            break
        data += 1
        skipped.append(data)
    first = data + 1
    firstcells = rows.get(first, [])
    key = firstcells[keycol - 1] if len(firstcells) >= keycol else None
    print(f'{scen} {tab:<28} heading r{hrow}  key col {keycol}  '
          f'skipped {skipped if skipped else "-"}  first data r{first} key={key!r}')
    for r in skipped:
        if any(str(v or "").strip().isdigit() and len(str(v).strip()) > 4 for v in rows[r]):
            bad.append(f'{scen}: line {r} looks like data but was skipped')

if bad:
    print('\n'.join(bad))
    sys.exit(1)
