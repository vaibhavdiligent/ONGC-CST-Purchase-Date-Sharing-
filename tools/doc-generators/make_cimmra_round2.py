"""Round two of the open points for Cipla - what is settled, and what is still needed.

Written after Cipla returned the first open-points document with four replies and
the vendor account group / number range workbook.
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()
st = doc.styles['Normal']
st.font.name = 'Calibri'; st.font.size = Pt(10.5)
for s in ('Heading 1', 'Heading 2', 'Heading 3'):
    doc.styles[s].font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    doc.styles[s].font.name = 'Calibri'
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Cm(2)
sec.left_margin = sec.right_margin = Cm(2.0)

def H(t, lvl=1): doc.add_heading(t, level=lvl)
def P(t='', bold=False, italic=False, size=None):
    p = doc.add_paragraph()
    r = p.add_run(t); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    return p

def T(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Grid Accent 1'
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ''
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(9.5)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            v = '' if v is None else str(v)
            cells[i].text = ''
            para = cells[i].paragraphs[0]
            for j, ln in enumerate(v.split('\n')):
                if j: para = cells[i].add_paragraph()
                rr = para.add_run(ln); rr.font.size = Pt(9.5)
    if widths:
        for r_ in t.rows:
            for i, w in enumerate(widths):
                r_.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t

# ------------------------------------------------------------------ title
doc.add_heading('Vendor Portal Interface - Business Partner Creation', level=0)
P('Open points - round two', bold=True, size=13)
doc.add_paragraph()

# ------------------------------------------------------------------ 1
H('1. Where we stand', 1)
P('Thank you for the four replies and for the account group and number range '
  'workbook. Between them they settle the Business Partner category and the number '
  'assignment, which were the two heaviest points.')
P('This document carries what is still needed: six follow-ups that arise from the '
  'replies themselves, the field mapping that has yet to be decided, two questions '
  'that were not answered in the last round, and the files and objects we need in '
  'order to build and test.')

# ------------------------------------------------------------------ 2
H('2. Settled - recorded so it is not asked again', 1)
T(['Point', 'Position'], [
    ['Business Partner category',
     'Organisation Business Partner for every account group, except Z008 (Employee '
     'Vendors) and Z014 (Employee Vendors - Associate), which are person Business '
     'Partners. For the Cimmra vendor file this means Z001 and Z002 are organisation '
     'Business Partners, proprietorships included.'],
    ['Number assignment',
     'External only for 0007, Z006, Z008, Z011, Z013 and Z014. Z001 and Z002 are '
     'internal. Vendor intervals: Z001 to Z1, Z002 to Z2, Z008 to 01, Z014 to A1.'],
    ['Business Partner number',
     'Same as the vendor number, captured in the existing SAPVendorN field. No new '
     'column is added to the response file.'],
    ['Previous account number (ALTKN)',
     'Optional, supplied by the portal, and only for Z008 and Z014.'],
], [4.5, 12.0])

# ------------------------------------------------------------------ 3
H('3. Follow-ups on the replies', 1)
P('Each of these arises directly from an answer given, and is needed before the '
  'answer can be implemented.', italic=True)
T(['#', 'What we need', 'Why'], [
    ['F1', 'The Business Partner grouping for each account group - the contents of '
     'CVIC_VEND_TO_BP1, and of CVIC_VEND_TO_BP2 for the roles.',
     'The workbook gives the vendor number ranges: every row carries object KREDITOR. '
     'A Business Partner takes its number from its grouping, which is a different '
     'number range object (BU_PARTNER). "Business Partner number = vendor number" '
     'holds only if, for each account group, the grouping points at the same interval '
     'and is flagged for the same number. We cannot confirm the answer to the number '
     'question until we can see that mapping.\n\n'
     'Worth noting with it: four vendor intervals are shared by two account groups '
     'each - Z1 by Z001 and Z015, Z3 by Y002 and Z012, 30 by Y001 and Z005, and XX by '
     '0007 and Z011. That is workable, but it means the account group to grouping '
     'mapping is not one to one and has to be read from the table rather than assumed.'],
    ['F2', 'Which column of which file carries the previous account number (ALTKN)?',
     'The vendor file has 124 columns and none of them can carry it. The reply says it '
     'comes from the portal for Z008 and Z014 - employee vendors, which arrive through '
     'ZFI_CIMMRA_EMPVEND, a file we have not seen. Please point us at the column, or '
     'send a sample of that file.'],
    ['F3', 'The contents of ZFI_VEN_GRP.',
     'Z001 and Z002 are confirmed as internally numbered, but today the program '
     'assigns the vendor number itself for any account group flagged in ZFI_VEN_GRP - '
     'it reads the highest existing vendor for the group and adds one. That is '
     'external behaviour on an internal group. The table will tell us whether this '
     'only ever fires for the six external groups, in which case there is no conflict '
     'and the logic simply disappears.'],
    ['F4', 'For Z008 and Z014, does the file carry first name and last name '
     'separately, or as one string?',
     'A person Business Partner stores a first name and a last name, not a single '
     'name. If the file sends one string we need the rule for splitting it, and a '
     'name such as "KURUBARABUDIHAL CHANDRASHEKAR CHINMAY SHETTY" cannot be split '
     'safely without one.'],
    ['F5', 'For the externally numbered groups, which column supplies the number, and '
     'is the Business Partner grouping external as well?',
     'With external assignment the number has to come from the file. With the '
     'Business Partner and the vendor carrying the same number, the grouping has to be '
     'external too, or the two cannot match.'],
    ['F6', 'Which account groups actually arrive through each interface file?',
     'The workbook lists 36. The Cimmra vendor file only ever carries Z001 and Z002 in '
     'the nine sample records; the employee files presumably carry Z008 and Z014. '
     'Confirming the set keeps us from building mapping for groups that never arrive.'],
], [1.2, 5.6, 9.7])

# ------------------------------------------------------------------ 4
H('4. Field mapping still to be decided', 1)
P('Cipla has asked that every column be mapped, whether or not the portal fills it. '
  'Of the 134 columns the program reads, 53 are posted and 81 are read and then '
  'discarded. Those 81 are the work, and they cannot be mapped until we know where '
  'they should go. We will propose a target for the obvious ones; the groups below '
  'need a decision.')
T(['#', 'Columns', 'What we need to know'], [
    ['M1', 'The location block - country, state, city, pin_code, address_line1 to 3, '
     'location_email, location_phone, location_type, sez_type, together with '
     'is_header_address_same_in_location',
     'Is this a second address for the vendor? A Business Partner holds several '
     'addresses, which the vendor master could not, so this can now be stored '
     'properly. Please confirm what the block represents and when the flag applies.'],
    ['M2', 'Tan, Cin, RegistrationNumber, TaxCode, TaxNo, w8_w9',
     'Which SAP field each of these should be written to.'],
    ['M3', 'gr_based_iv, schema_code',
     'The program hard-codes GR based invoice verification as X, and derives the '
     'schema group from a set rather than from the file. Mapping these columns '
     'changes current behaviour - please confirm that is wanted.'],
    ['M4', 'bnkcountry',
     'The vendor’s own country is used as the bank country today. Mapping this column '
     'changes behaviour for foreign vendors.'],
    ['M5', 'bank_name, branch_address, bic, micr, iban',
     'Which bank master field each belongs to. Two points to settle with it: a '
     '24-character IBAN is being written into the 18-character bank account number '
     'and truncated, while the file’s own iban column is discarded; and bic holds a '
     'branch name in eight of the nine sample records, not a BIC.'],
    ['M6', 'contact_number, contact_purpose, Contactlocation_id, '
     'SupplierIntroducerName / Id / Email, CompanyWebsite, Phone',
     'Should these become a Business Partner contact person, or communication data on '
     'the address?'],
    ['M7', 'Status, is_agree, IsUpdateBySupplier, CreatedSource, SupplierInvitationId, '
     'and the GST and PAN verification flags',
     'These look like portal control fields. Should they be stored, and if so where, '
     'or ignored?'],
    ['M8', 'Address3',
     'Populated but not used today. We propose street supplement 2.'],
    ['M9', 'CountryName, company, payment_term, purchasing_organization, reco_account, '
     'industry_key, schema, currency, gst_classification, OrganizationTypeCode',
     'We assume these are text echoes of the code columns next to them and that '
     'nothing is stored. Please confirm.'],
], [1.2, 5.6, 9.7])

# ------------------------------------------------------------------ 5
H('5. Two questions still open from the last round', 1)
T(['#', 'Question', 'Why we are asking'], [
    ['Q1', 'Please provide the mapping from the portal’s State code to the SAP region.',
     'The portal’s codes are self-consistent - Gujarat 06, Karnataka 10, Madhya '
     'Pradesh 12, Maharashtra 13 - but they are not the codes carried in the GST '
     'numbers on the same records (24, 29, 23, 27), and two records fit neither '
     'pattern (Kolkata 25, Hyderabad 01). The value is written straight into the '
     'address region today, so it is worth confirming for the existing interface as '
     'well.'],
    ['Q2', 'Do the columns always appear in the same order, and is a column the portal '
     'does not send always absent from the end of the line rather than from the middle?',
     'The line is read by position, not by heading. A column dropped from the middle '
     'shifts every value after it, and the record is created with the wrong value in '
     'every remaining field - silently, with no error. The five sample files carry 124 '
     'columns while the program reads 134, so the last ten are already absent; that is '
     'handled, but only because they are at the end.'],
], [1.2, 5.6, 9.7])

# ------------------------------------------------------------------ 6
H('6. Files and objects we need', 1)
T(['#', 'What', 'Why'], [
    ['R1', 'One sample file each for vendor change, employee vendor creation, employee '
     'vendor change and vendor extension, and one record carrying withholding tax and '
     'exemption data',
     'All nine records supplied are creations with no withholding tax, so those paths '
     'cannot be tested against a real file. The employee file also answers F2 and F4.'],
    ['R2', 'One sample response file',
     'To locate the SAPVendorN field in the layout we have, and confirm nothing else '
     'has to move.'],
    ['R3', 'ZGEN_UPDATE_X and ZPARAM_TABLE',
     'Used throughout. We have the calls but not the definitions.'],
    ['R4', 'The ZKRP_* structures used by the two function modules',
     'About forty structures carrying the interface parameters.'],
    ['R5', 'ZFI_VEND_MSME and ZFI_TAX_TAB - the table definitions',
     'Both are written directly by the interface. We need the field types before we '
     'can map to them.'],
    ['R6', 'The screen enhancement on XK01 that maintains ZFI_TAX_TAB',
     'Needed for the vendor extension program, where the table is filled from the '
     'screen rather than from the file.'],
    ['R7', 'A development system and a transport, and development authorisation for '
     'vendor and Business Partner master data, with read access to the AL11 inbound '
     'and outbound directories',
     'To develop and to test against the real files.'],
], [1.2, 5.6, 9.7])

import os
out = os.path.join(os.path.dirname(__file__), '..', '..', 'docs', 'cimmra',
                   'Vendor_BP_Open_Points_Round2.docx')
doc.save(os.path.abspath(out))
print('written', os.path.abspath(out))
