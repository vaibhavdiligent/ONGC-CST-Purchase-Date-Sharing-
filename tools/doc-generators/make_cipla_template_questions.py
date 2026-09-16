"""The three open questions on the customer code template workbook."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()
st = doc.styles['Normal']
st.font.name = 'Calibri'; st.font.size = Pt(10.5)
for s in ('Heading 1', 'Heading 2'):
    doc.styles[s].font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    doc.styles[s].font.name = 'Calibri'
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Cm(2)
sec.left_margin = sec.right_margin = Cm(2.0)

def P(t='', bold=False, italic=False, size=None):
    p = doc.add_paragraph()
    r = p.add_run(t); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    return p

def T(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Grid Accent 1'
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ''
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(9.5)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            para = cells[i].paragraphs[0]
            for j, ln in enumerate(str(v).split('\n')):
                if j: para = cells[i].add_paragraph()
                rr = para.add_run(ln); rr.font.size = Pt(9.5)
    for r_ in t.rows:
        for i, w in enumerate(widths):
            r_.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t

doc.add_heading('Customer Code Templates - three points to confirm', level=0)
P('We have worked through the template workbook and the programs already running, and '
  'everything else is resolved. Three points cannot be settled from the workbook itself.',
  italic=True)
doc.add_paragraph()

T(['#', 'Question', 'Why we cannot answer it ourselves'], [
    ['1',
     'The Europe template - which countries does it serve?\n\n'
     'And do company codes 7450 and 7000 use the same template, or different ones?',
     'Every other sheet names its country, in the sheet name and in the Project row above '
     'the block - Australia, Dubai, India, Kenya, Morocco, Uganda, and "US Sold to’s" '
     'for Exelan and Invagen. Europe is a region rather than a country, and the sheet '
     'carries two company codes. The download is selected by country, so we need the '
     'country keys the Europe template covers.'],
    ['2',
     'The QCIL export template - is it the 83-column one or the 65-column one?',
     'That block carries two layouts at once. Its technical field row is complete and '
     'coherent: 83 columns, identical to the template Australia and Morocco use. Its '
     'description row has only 65 entries, drifts out of step from column 2 onward, and '
     'stops before the licence and bank-guarantee columns - and the sample data row '
     'follows the description row, not the technical row. Two rows agree with each other '
     'and the third is the complete one, so this is two templates pasted together rather '
     'than one row being out of line.'],
    ['3',
     'In the India ZSHM template (6 columns) and the cust extn template (5 columns), '
     'which tax category does each "Tax classification for customer" column stand for?',
     'Everywhere else the workbook names the category in the heading - JOCG, JTC1, JTX1 '
     'to JTX4 on the India sheets, UTXJ, UTX2 and UTX3 on the US sheets. In these two '
     'templates every one of those columns is headed only "Tax classification for '
     'customer". The order is the only clue, and a tax classification written against '
     'the wrong category is not visible on the screen afterwards.'],
], [1.0, 6.4, 9.1])

P('Nothing else is outstanding. The country for every other sheet, the field behind every '
  'column including the licence, bank-guarantee and Aadhaar data, and the template for '
  'every other country and account group combination have all been resolved from the '
  'workbook and from the programs already in use.', italic=True)

import os
out = os.path.join(os.path.dirname(__file__), '..', '..', 'docs', 'cipla',
                   'Customer_Code_Templates_Points_to_Confirm.docx')
doc.save(os.path.abspath(out))
print('written', os.path.abspath(out))
